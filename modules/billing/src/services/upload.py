"""SP-1 Plan B Task 2 — Upload service.

Sub-functions:
  - compute_sha256, write_upload_file (atomic via .tmp -> os.replace)
  - find_existing_upload (sha256 dedup, tenant-scoped)
  - parse_csv_bytes, parse_xlsx_bytes (parsers)
  - validate_row (per-row business rules; PHI-safe error messages)
  - parse_upload (orchestrator)
  - supersede_upload (immutable provenance)

Financial precision (.claude/rules/financial-precision.md):
  - Decimal(str(value)) at parse; never Decimal(float)
  - amount_billed max 4dp, quantity max 3dp (matches HEAD claim_records)

PHI controls (.claude/rules/phi-compliance.md):
  - row_errors NEVER echo member_id values; only describe failure category
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import uuid
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from src.models.tables import ClaimRecord, Upload, UploadStatus

# LESSON-004 anchors: \A...\Z reject trailing newlines that ^...$ would accept
_NDC_RE = re.compile(r"\A\d{11}\Z")
_NPI_RE = re.compile(r"\A\d{10}\Z")

_REQUIRED_COLUMNS = (
    "ndc", "npi", "claim_id", "date_of_service",
    "quantity", "days_supply", "amount_billed", "member_id",
)
_MAX_MEMBER_ID_LEN = 64
_QUANTITY_MAX_DP = 3
_AMOUNT_MAX_DP = 4


def compute_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def write_upload_file(
    *,
    base_dir: Path,
    tenant_id: uuid.UUID,
    upload_id: uuid.UUID,
    filename: str,
    content: bytes,
) -> Path:
    # P1-security: strip all directory components so an attacker-supplied
    # filename like "../../etc/passwd" or an absolute Windows path cannot
    # escape the tenant/upload directory.  Path.name gives only the
    # final component; fall back to "upload.bin" for empty results.
    safe_name = Path(filename).name or "upload.bin"
    target_dir = base_dir / str(tenant_id) / str(upload_id)
    target_dir.mkdir(parents=True, exist_ok=True)
    final = target_dir / safe_name
    tmp = final.with_suffix(final.suffix + ".tmp")
    tmp.write_bytes(content)
    os.replace(tmp, final)
    return final


def find_existing_upload(
    session: Session, *, tenant_id: uuid.UUID, sha256: str
) -> Upload | None:
    stmt = select(Upload).where(
        Upload.tenant_id == tenant_id, Upload.sha256 == sha256
    )
    return session.execute(stmt).scalar_one_or_none()


def _luhn_check(digits: str) -> bool:
    total = 0
    for i, ch in enumerate(reversed(digits)):
        d = int(ch)
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def _validate_ndc(value: str) -> str | None:
    if not _NDC_RE.match(value):
        return "ndc must be exactly 11 digits"
    return None


def _validate_npi(value: str) -> str | None:
    if not _NPI_RE.match(value):
        return "npi must be exactly 10 digits"
    # NPI Luhn check uses the 80840 prefix per CMS spec
    # (.claude/rules/security.md "MUST validate NPI with Luhn check (prefix 80840)")
    if not _luhn_check("80840" + value):
        return "npi fails Luhn check"
    return None


def _validate_decimal(
    value: str, *, field: str, max_dp: int, allow_zero: bool = True
) -> tuple[Decimal | None, str | None]:
    try:
        d = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None, f"{field} must be a decimal number"
    if d < 0:
        return None, f"{field} must be >= 0"
    if not allow_zero and d == 0:
        return None, f"{field} must be > 0"
    exp = d.as_tuple().exponent
    if isinstance(exp, int) and exp < -max_dp:
        return None, f"{field} has more than {max_dp} decimal places"
    return d, None


def validate_row(row: dict[str, str]) -> tuple[dict[str, Any] | None, str | None]:
    err = _validate_ndc(row.get("ndc", ""))
    if err:
        return None, err
    err = _validate_npi(row.get("npi", ""))
    if err:
        return None, err

    try:
        dos = date.fromisoformat(row.get("date_of_service", ""))
    except (ValueError, TypeError):
        return None, "date_of_service must be YYYY-MM-DD"

    qty_dec, err = _validate_decimal(
        row.get("quantity", ""), field="quantity", max_dp=_QUANTITY_MAX_DP,
    )
    if err:
        return None, err

    try:
        days = int(row.get("days_supply", ""))
    except (ValueError, TypeError):
        return None, "days_supply must be a positive integer"
    if days < 1:
        return None, "days_supply must be >= 1"

    amount_dec, err = _validate_decimal(
        row.get("amount_billed", ""), field="amount_billed", max_dp=_AMOUNT_MAX_DP,
    )
    if err:
        return None, err

    member_id = row.get("member_id", "")
    if not isinstance(member_id, str) or member_id == "":
        return None, "member_id must be non-empty"
    if len(member_id) > _MAX_MEMBER_ID_LEN:
        return None, f"member_id exceeds {_MAX_MEMBER_ID_LEN} character limit"

    claim_id = row.get("claim_id", "")
    if not claim_id:
        return None, "claim_id must be non-empty"

    return {
        "ndc": row["ndc"],
        "npi": row["npi"],
        "claim_id": claim_id,
        "date_of_service": dos,
        "quantity": qty_dec,
        "days_supply": days,
        "amount_billed": amount_dec,
        "member_id": member_id,
    }, None


def _check_required_columns(headers: list[str]) -> None:
    headers_set = {h.strip().lower() for h in headers}
    missing = set(_REQUIRED_COLUMNS) - headers_set
    if missing:
        raise ValueError(f"CSV missing required columns: {sorted(missing)}")


def parse_csv_bytes(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8")
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise ValueError("CSV missing header row")
    _check_required_columns(list(reader.fieldnames))
    return [dict(r) for r in reader]


def parse_xlsx_bytes(content: bytes) -> list[dict[str, str]]:
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to parse .xlsx uploads") from exc
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows, [])]
    _check_required_columns(headers)
    out: list[dict[str, str]] = []
    for row in rows:
        if all(c is None for c in row):
            continue
        out.append({h: ("" if v is None else str(v)) for h, v in zip(headers, row, strict=False)})
    return out


@dataclass(frozen=True)
class UploadParseResult:
    upload: Upload
    claims_written: int


def parse_upload(
    session: Session, *, upload: Upload, file_bytes: bytes
) -> UploadParseResult:
    # P2-xlsx: dispatch to the correct parser based on mime_type/filename.
    # Callers that pass raw bytes from an xlsx file must set upload.mime_type
    # to "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # or upload.filename ending in ".xlsx".
    _xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    _is_xlsx = (upload.mime_type or "").startswith(_xlsx_mime) or (
        upload.filename or ""
    ).lower().endswith(".xlsx")
    rows = parse_xlsx_bytes(file_bytes) if _is_xlsx else parse_csv_bytes(file_bytes)
    row_count = len(rows)
    errors: list[dict[str, Any]] = []
    valid_rows: list[dict[str, Any]] = []

    for idx, raw in enumerate(rows, start=1):
        parsed, err = validate_row(raw)
        if err is not None:
            errors.append({"row": idx, "message": err})
        else:
            assert parsed is not None
            valid_rows.append(parsed)

    if row_count > 0 and len(valid_rows) == 0:
        upload.status = UploadStatus.validation_failed.value
        upload.row_count = row_count
        upload.error_count = len(errors)
        upload.row_errors = errors
        session.flush()
        return UploadParseResult(upload=upload, claims_written=0)

    now = datetime.now(UTC)
    for row in valid_rows:
        claim = ClaimRecord(
            id=uuid.uuid4(),
            tenant_id=upload.tenant_id,
            source_type="upload",
            auth_number=row["claim_id"],
            claim_type="rx",
            pharmacy_npi=row["npi"],
            ndc=row["ndc"],
            quantity=row["quantity"],
            days_supply=row["days_supply"],
            date_of_service=row["date_of_service"],
            date_received=now,
            member_id=row["member_id"],
            net_amount=Decimal("0"),
            created_at=now,
            upload_id=upload.id,
            amount_billed=row["amount_billed"],
        )
        session.add(claim)

    upload.status = UploadStatus.validated.value
    upload.row_count = row_count
    upload.error_count = len(errors)
    upload.row_errors = errors if errors else None
    session.flush()
    return UploadParseResult(upload=upload, claims_written=len(valid_rows))


def supersede_upload(
    session: Session, *, old_upload: Upload, new_upload: Upload
) -> None:
    if old_upload.tenant_id != new_upload.tenant_id:
        raise ValueError("cannot supersede across tenants")
    # P2-supersede: delete ClaimRecord rows from the old upload before
    # parse_upload inserts replacements.  The unique constraint
    # uq_claim_tenant_auth (tenant_id, auth_number) fires when the replacement
    # file re-uses the same claim_id values — there is no way to retain old
    # rows with the same (tenant_id, auth_number) while inserting new ones.
    # Provenance immutability covers the Upload row (status → superseded,
    # supersedes_upload_id chain preserved); individual ClaimRecord rows are
    # processing artifacts replaced in full by the superseding upload.
    old_claims = session.execute(
        select(ClaimRecord).where(
            ClaimRecord.upload_id == old_upload.id,
            ClaimRecord.tenant_id == old_upload.tenant_id,
        )
    ).scalars().all()
    for c in old_claims:
        session.delete(c)
    session.flush()
    old_upload.status = UploadStatus.superseded.value
    new_upload.supersedes_upload_id = old_upload.id
    session.flush()
