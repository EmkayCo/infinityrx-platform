"""Audit query + export API tests."""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from src.audit.api import build_audit_router
from src.audit.models import AuditLog
from src.audit.schemas import AuditEntry
from src.audit.service import AuditService

TENANT = uuid.UUID("11111111-1111-1111-1111-111111111111")
OTHER = uuid.UUID("22222222-2222-2222-2222-222222222222")
USER = uuid.UUID("aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa")


@dataclass
class _User:
    id: uuid.UUID
    tenant_id: uuid.UUID
    permissions: set[str]


def _require_permission(perm: str):
    def dep():
        user = _User(id=USER, tenant_id=TENANT, permissions={"audit:read", "audit:export"})
        if perm not in user.permissions:
            raise HTTPException(status_code=403, detail={"error": "forbidden"})
        return user

    return dep


@pytest.fixture
def app(db_session):
    db_session.execute(AuditLog.__table__.delete())
    db_session.commit()

    def _get_session():
        yield db_session

    app = FastAPI()
    app.include_router(
        build_audit_router(
            get_session=_get_session,
            require_permission=_require_permission,
        )
    )
    return app


def _seed(db_session, *, tenant=TENANT, count=3):
    svc = AuditService(db_session)
    for i in range(count):
        svc.log(
            AuditEntry(
                tenant_id=tenant,
                user_id=USER,
                action="create",
                module="core-platform",
                entity_type="user",
                entity_id=f"u{i}",
                after_value={"i": i},
            )
        )
    db_session.commit()


def test_list_returns_page(app, db_session):
    _seed(db_session)
    client = TestClient(app)
    resp = client.get("/api/v1/audit")
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 3
    assert len(body["items"]) == 3


def test_list_filter_by_action(app, db_session):
    _seed(db_session)
    client = TestClient(app)
    resp = client.get("/api/v1/audit", params={"action": "create"})
    assert resp.json()["total"] == 3
    resp = client.get("/api/v1/audit", params={"action": "delete"})
    assert resp.json()["total"] == 0


def test_list_tenant_isolation(app, db_session):
    _seed(db_session, tenant=TENANT, count=2)
    _seed(db_session, tenant=OTHER, count=5)
    client = TestClient(app)
    resp = client.get("/api/v1/audit")
    assert resp.json()["total"] == 2


def test_get_single_and_404(app, db_session):
    _seed(db_session)
    client = TestClient(app)
    resp = client.get("/api/v1/audit")
    first_id = resp.json()["items"][0]["id"]
    got = client.get(f"/api/v1/audit/{first_id}")
    assert got.status_code == 200
    missing = client.get("/api/v1/audit/9999")
    assert missing.status_code == 404


def test_export_csv(app, db_session):
    _seed(db_session)
    client = TestClient(app)
    resp = client.get("/api/v1/audit/export", params={"format": "csv"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    text = resp.text
    assert "entity_id" in text.splitlines()[0]
    assert "u0" in text


def test_export_xlsx(app, db_session):
    _seed(db_session)
    client = TestClient(app)
    resp = client.get("/api/v1/audit/export", params={"format": "xlsx"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["audit"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 4  # header + 3 seed rows


def test_export_rejects_bad_format(app):
    client = TestClient(app)
    resp = client.get("/api/v1/audit/export", params={"format": "pdf"})
    assert resp.status_code == 422
