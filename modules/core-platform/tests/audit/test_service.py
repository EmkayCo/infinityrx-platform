"""AuditService unit tests: writes, filter queries, tenant isolation, exports."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook
from src.audit.schemas import AuditEntry, AuditQuery
from src.audit.service import EXPORT_COLUMNS, AuditService


def _entry(tenant: uuid.UUID, **overrides) -> AuditEntry:
    base = dict(
        tenant_id=tenant,
        user_id=uuid.uuid4(),
        action="create",
        module="core-platform",
        entity_type="user",
        entity_id="u1",
        before_value=None,
        after_value={"name": "Alice"},
        ip_address="127.0.0.1",
        user_agent="pytest",
        correlation_id=uuid.uuid4(),
    )
    base.update(overrides)
    return AuditEntry(**base)


def test_log_persists_entry(db_session, tenant_id):
    svc = AuditService(db_session)
    row = svc.log(_entry(tenant_id))
    db_session.commit()
    assert row.id is not None
    assert row.tenant_id == str(tenant_id)


def test_query_filters_by_each_field(db_session, tenant_id, user_id):
    svc = AuditService(db_session)
    cid = uuid.uuid4()
    svc.log(_entry(tenant_id, action="create", user_id=user_id, correlation_id=cid))
    svc.log(_entry(tenant_id, action="delete", module="billing", entity_type="batch", entity_id="b9"))
    db_session.commit()

    # action
    page = svc.query(tenant_id, AuditQuery(action="create"))
    assert page.total == 1 and page.items[0].action == "create"
    # module
    page = svc.query(tenant_id, AuditQuery(module="billing"))
    assert page.total == 1 and page.items[0].module == "billing"
    # entity_type
    page = svc.query(tenant_id, AuditQuery(entity_type="batch"))
    assert page.total == 1
    # entity_id
    page = svc.query(tenant_id, AuditQuery(entity_id="b9"))
    assert page.total == 1
    # user_id
    page = svc.query(tenant_id, AuditQuery(user_id=user_id))
    assert page.total == 1
    # correlation_id
    page = svc.query(tenant_id, AuditQuery(correlation_id=cid))
    assert page.total == 1
    # date range
    page = svc.query(
        tenant_id,
        AuditQuery(date_from=datetime.now(UTC) - timedelta(hours=1), date_to=datetime.now(UTC) + timedelta(hours=1)),
    )
    assert page.total == 2


def test_tenant_isolation_blocks_cross_tenant_reads(db_session, tenant_id, other_tenant_id):
    svc = AuditService(db_session)
    svc.log(_entry(tenant_id))
    svc.log(_entry(other_tenant_id))
    db_session.commit()

    page = svc.query(tenant_id, AuditQuery())
    assert page.total == 1
    assert str(page.items[0].tenant_id) == str(tenant_id)

    # Even a crafted filter cannot leak another tenant's row
    crafted = AuditQuery(entity_id="u1")
    page = svc.query(tenant_id, crafted)
    assert all(str(i.tenant_id) == str(tenant_id) for i in page.items)


def test_pagination(db_session, tenant_id):
    svc = AuditService(db_session)
    for i in range(5):
        svc.log(_entry(tenant_id, entity_id=f"u{i}"))
    db_session.commit()
    page1 = svc.query(tenant_id, AuditQuery(limit=2, offset=0))
    page2 = svc.query(tenant_id, AuditQuery(limit=2, offset=2))
    assert page1.total == 5 and page2.total == 5
    assert {i.entity_id for i in page1.items}.isdisjoint({i.entity_id for i in page2.items})


def test_get_single_entry_and_missing(db_session, tenant_id, other_tenant_id):
    svc = AuditService(db_session)
    row = svc.log(_entry(tenant_id))
    db_session.commit()
    fetched = svc.get(tenant_id, row.id)
    assert fetched is not None and fetched.id == row.id
    # Cross-tenant must return None
    assert svc.get(other_tenant_id, row.id) is None
    # Missing id
    assert svc.get(tenant_id, 999999) is None


def test_export_csv_contains_header_and_rows(db_session, tenant_id):
    svc = AuditService(db_session)
    svc.log(_entry(tenant_id, entity_id="u-csv"))
    db_session.commit()
    chunks = list(svc.export_csv(tenant_id, AuditQuery()))
    text = b"".join(chunks).decode()
    assert text.splitlines()[0] == ",".join(EXPORT_COLUMNS)
    assert "u-csv" in text


def test_export_xlsx_is_valid_workbook(db_session, tenant_id):
    svc = AuditService(db_session)
    svc.log(_entry(tenant_id, entity_id="u-xlsx"))
    db_session.commit()
    content = svc.export_xlsx(tenant_id, AuditQuery())
    wb = load_workbook(io.BytesIO(content))
    ws = wb["audit"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == tuple(EXPORT_COLUMNS)
    assert any("u-xlsx" in str(cell) for cell in rows[1])


def test_iter_rows_respects_max_rows(db_session, tenant_id):
    svc = AuditService(db_session)
    for i in range(3):
        svc.log(_entry(tenant_id, entity_id=f"u{i}"))
    db_session.commit()
    rows = list(svc.iter_rows(tenant_id, AuditQuery(), max_rows=2))
    assert len(rows) == 2


def test_row_to_read_handles_nullables(db_session, tenant_id):
    svc = AuditService(db_session)
    svc.log(_entry(tenant_id, user_id=None, correlation_id=None))
    db_session.commit()
    page = svc.query(tenant_id, AuditQuery())
    assert page.items[0].user_id is None
    assert page.items[0].correlation_id is None
