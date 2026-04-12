"""AuditService: thin persistence layer over SQLAlchemy.

Used by the middleware, the PHI decorator, and the audit query API.
Every operation is tenant-scoped; callers must supply ``tenant_id`` and
the query helpers enforce it server-side to defeat parameter tampering.
"""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import Iterator
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from src.audit.models import AuditLog, _as_uuid_str
from src.audit.schemas import AuditEntry, AuditEntryRead, AuditPage, AuditQuery

EXPORT_COLUMNS = [
    "id",
    "created_at",
    "tenant_id",
    "user_id",
    "action",
    "module",
    "entity_type",
    "entity_id",
    "ip_address",
    "user_agent",
    "correlation_id",
]


class AuditService:
    """Persistence + query facade for core.audit_log."""

    def __init__(self, session: Session) -> None:
        self._session = session

    # ------------------------------------------------------------------
    # Writes
    # ------------------------------------------------------------------
    def log(self, entry: AuditEntry) -> AuditLog:
        row = AuditLog(
            tenant_id=str(entry.tenant_id),
            user_id=_as_uuid_str(entry.user_id),
            action=entry.action,
            module=entry.module,
            entity_type=entry.entity_type,
            entity_id=entry.entity_id,
            before_value=entry.before_value,
            after_value=entry.after_value,
            ip_address=entry.ip_address,
            user_agent=entry.user_agent,
            correlation_id=_as_uuid_str(entry.correlation_id),
        )
        self._session.add(row)
        self._session.flush()
        return row

    # ------------------------------------------------------------------
    # Queries (tenant-scoped — callers MUST pass the viewer's tenant_id)
    # ------------------------------------------------------------------
    def _base_stmt(self, tenant_id: uuid.UUID, q: AuditQuery):
        stmt = select(AuditLog).where(AuditLog.tenant_id == str(tenant_id))
        if q.action is not None:
            stmt = stmt.where(AuditLog.action == q.action)
        if q.module is not None:
            stmt = stmt.where(AuditLog.module == q.module)
        if q.entity_type is not None:
            stmt = stmt.where(AuditLog.entity_type == q.entity_type)
        if q.entity_id is not None:
            stmt = stmt.where(AuditLog.entity_id == q.entity_id)
        if q.user_id is not None:
            stmt = stmt.where(AuditLog.user_id == str(q.user_id))
        if q.correlation_id is not None:
            stmt = stmt.where(AuditLog.correlation_id == str(q.correlation_id))
        if q.date_from is not None:
            stmt = stmt.where(AuditLog.created_at >= q.date_from)
        if q.date_to is not None:
            stmt = stmt.where(AuditLog.created_at <= q.date_to)
        return stmt

    def query(self, tenant_id: uuid.UUID, q: AuditQuery) -> AuditPage:
        stmt = self._base_stmt(tenant_id, q).order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        total_stmt = select(func.count()).select_from(self._base_stmt(tenant_id, q).subquery())
        total = int(self._session.execute(total_stmt).scalar_one())
        rows = (
            self._session.execute(stmt.limit(q.limit).offset(q.offset)).scalars().all()
        )
        items = [self._row_to_read(r) for r in rows]
        return AuditPage(items=items, total=total, limit=q.limit, offset=q.offset)

    def get(self, tenant_id: uuid.UUID, audit_id: int) -> AuditEntryRead | None:
        row = self._session.execute(
            select(AuditLog).where(
                AuditLog.id == audit_id,
                AuditLog.tenant_id == str(tenant_id),
            )
        ).scalar_one_or_none()
        return self._row_to_read(row) if row is not None else None

    # ------------------------------------------------------------------
    # Export (streaming)
    # ------------------------------------------------------------------
    def iter_rows(
        self, tenant_id: uuid.UUID, q: AuditQuery, *, max_rows: int = 1_000_000
    ) -> Iterator[AuditLog]:
        stmt = (
            self._base_stmt(tenant_id, q)
            .order_by(AuditLog.id.asc())
            .limit(max_rows)
        )
        yield from self._session.execute(stmt).scalars()

    def export_csv(
        self, tenant_id: uuid.UUID, q: AuditQuery, *, max_rows: int = 1_000_000
    ) -> Iterator[bytes]:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(EXPORT_COLUMNS)
        yield buffer.getvalue().encode()
        buffer.seek(0)
        buffer.truncate()
        for row in self.iter_rows(tenant_id, q, max_rows=max_rows):
            writer.writerow([self._field(row, c) for c in EXPORT_COLUMNS])
            yield buffer.getvalue().encode()
            buffer.seek(0)
            buffer.truncate()

    def export_xlsx(
        self, tenant_id: uuid.UUID, q: AuditQuery, *, max_rows: int = 1_000_000
    ) -> bytes:
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("audit")
        ws.append(EXPORT_COLUMNS)
        for row in self.iter_rows(tenant_id, q, max_rows=max_rows):
            ws.append([self._field(row, c) for c in EXPORT_COLUMNS])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------
    @staticmethod
    def _field(row: AuditLog, column: str):
        value = getattr(row, column)
        if isinstance(value, datetime):
            return value.isoformat()
        return value

    @staticmethod
    def _row_to_read(row: AuditLog) -> AuditEntryRead:
        return AuditEntryRead(
            id=row.id,
            tenant_id=uuid.UUID(row.tenant_id),
            user_id=uuid.UUID(row.user_id) if row.user_id else None,
            action=row.action,
            module=row.module,
            entity_type=row.entity_type,
            entity_id=row.entity_id,
            before_value=row.before_value,
            after_value=row.after_value,
            ip_address=row.ip_address,
            user_agent=row.user_agent,
            correlation_id=uuid.UUID(row.correlation_id) if row.correlation_id else None,
            created_at=row.created_at,
        )
